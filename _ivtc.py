from copy import deepcopy
from datetime import datetime
from pathlib import Path
import random

from docx import Document
import openpyxl as op


class ExcelToTest:

    DB_SHEET_NAME: str = "DB"

    ANIMALS = ["unicorns", "dragons", "ponies",
               "griffins", "hippos", "otters", "bears"]

    LETTERS = "ABCDEFGHIJKLMNOPQRSTUVWXYZ"

    db_file: Path
    variant_type: str
    variant_cnt: int
    verb_cnt: int
    output_dir: Path

    def __init__(self,
                 db_file: Path,
                 variant_type: str,
                 variant_cnt: int,
                 verb_cnt: int,
                 output_dir: Path):

        self.time_of_gen = datetime.now()

        self.db_file = db_file
        assert db_file.is_file()

        self.variant_type = variant_type
        assert variant_type in ["Animals", "Numbers", "Letters"]

        self.variant_cnt = variant_cnt
        if variant_type == "Animals":
            assert variant_cnt <= len(self.ANIMALS)

        if variant_type == "Letters":
            assert variant_cnt <= len(self.LETTERS)

        self.verb_cnt = verb_cnt

        self.debug = False

        self.output_dir = output_dir
        assert output_dir.is_dir()

    def enable_debugging(self):
        self.debug = True

    def _l(self, msg):
        if self.debug:
            print(msg)

    def _get_group_name(self, i: int) -> str:
        if self.variant_type == "Animals":
            return self.ANIMALS[i]

        elif self.variant_type == "Letters":
            return self.LETTERS[i]

        elif self.variant_type == "Numbers":
            return str(i + 1)

        else:
            assert False

    def _load_xslx(self):
        return op.load_workbook(str(self.db_file))

    def _check_format(self, db_sheet):
        assert db_sheet["A1"].value == "Base Form"
        assert db_sheet["B1"].value == "Past Simple"
        assert db_sheet["C1"].value == "Past Participle"
        assert db_sheet["D1"].value == "Czech Translation"

    @staticmethod
    def clean_str(dirty_str: str) -> str:
        if dirty_str is not None:
            dirty_str = dirty_str.replace("\u00A0", " ")
            dirty_str = dirty_str.strip()
        return dirty_str

    def get_time_str(self):
        dt = self.time_of_gen.strftime("%d. %m. %Y")
        file_dt = self.time_of_gen.strftime("%Y%m%d%H%M%S")

        return (dt, file_dt)

    def compose_a_doc(self, grp_name: str,
                      is_teacher: bool,
                      verbs: list[dict]):
        grp_name = grp_name.capitalize()

        doc = Document()
        category = "Student's Test"
        file_category = "test"
        if is_teacher:
            category = "Teacher's Key"
            file_category = "key"

        dt, file_dt = self.get_time_str()

        doc.add_heading(
            f'{grp_name}: a {category} ___ / {self.verb_cnt*3} pts    Name: ________________', 2)
        doc.add_paragraph(
            'Fill in the empty cells with correct versions and/or translations of the correct irregular verbs.')

        table = doc.add_table(rows=1, cols=4)
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Basic Form'
        hdr_cells[1].text = 'Past Simple'
        hdr_cells[2].text = 'Past Participle'
        hdr_cells[3].text = 'Czech Translation'
        for a_verb in verbs:
            row_cells = table.add_row().cells
            row_cells[0].text = a_verb["base"]
            row_cells[1].text = a_verb["past_simple"]
            row_cells[2].text = a_verb["past_participle"] if a_verb["past_participle"] is not None else ""
            row_cells[3].text = a_verb["cz"]

        doc.add_page_break()

        doc.save(str(self.output_dir /
                     f"{grp_name}_{file_category}_{file_dt}.docx"))

    def compose_docx(self, variants: list[dict]):

        for a_group in variants:
            self._l("Composing teacher's key...")
            self.compose_a_doc(a_group, True, variants[a_group]["teacher"])

            self._l("Composing student's test...")
            self.compose_a_doc(a_group, False, variants[a_group]["student"])

    @staticmethod
    def studentize(teachers_verbs: list[dict]) -> dict:

        students_verbs = deepcopy(teachers_verbs)

        idxs = ["base", "past_simple", "past_participle", "cz"]
        total_cnt = len(idxs)

        for a_verb in students_verbs:
            choice = random.randint(0, total_cnt - 1)
            if a_verb[idxs[choice]] is None:
                choice = (choice + 1) % total_cnt
                assert a_verb[idxs[choice]] is not None

            for idx in idxs:
                if idx != idxs[choice]:
                    a_verb[idx] = ""

        return students_verbs

    def xform(self):
        wb = self._load_xslx()

        db_sheet = wb[self.DB_SHEET_NAME]
        self._check_format(db_sheet)

        row_cnt = 0

        all_verbs = []

        for row in db_sheet.iter_rows():
            if row_cnt != 0:
                a_verb = {
                    "base": self.clean_str(row[0].value),
                    "past_simple": self.clean_str(row[1].value),
                    "past_participle": self.clean_str(row[2].value),
                    "cz": self.clean_str(row[3].value)
                }
                all_verbs.append(a_verb)
            row_cnt += 1

        self._l(f"Parsed {row_cnt} rows!")
        assert row_cnt >= self.verb_cnt

        # self._l(json.dumps(all_verbs, indent=2))

        variants = {}

        for i in range(self.variant_cnt):
            group_name = self._get_group_name(i)
            self._l(f"Generating variant {group_name}...")

            verb_selection = random.sample(all_verbs, self.verb_cnt)
            # self._l(json.dumps(verb_selection, indent=2))

            random.shuffle(verb_selection)
            # self._l(json.dumps(verb_selection, indent=2))
            variants[group_name] = {}
            variants[group_name]["teacher"] = verb_selection
            variants[group_name]["student"] = self.studentize(verb_selection)

        self._l(f"{variants}")
        self.compose_docx(variants)
